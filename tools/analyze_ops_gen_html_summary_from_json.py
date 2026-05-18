#!/usr/bin/env python3
"""
算子测试结果分析工具 (v3 - JSON汇总 + HTML报告 + Excel生成)
用法: python analyze_ops_gen_html_summary_from_json.py <结果文件夹路径>
输出: 在目标文件夹下生成 report.html 和 summary.xlsx
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl")
    print("请运行: pip install openpyxl")
    sys.exit(1)

try:
    import numpy as np
except ImportError:
    np = None


def load_results(folder):
    """加载所有 summary*.json 文件并合并结果"""
    merged = {}
    env = {}
    timestamp = ""

    json_files = sorted(folder.glob("summary*.json"))
    if not json_files:
        print("错误: 未找到 summary*.json 文件")
        sys.exit(1)

    for jf in json_files:
        print(f"  读取: {jf.name}")
        with open(jf, "r", encoding="utf-8") as f:
            data = json.load(f)

        if not timestamp and "timestamp" in data:
            timestamp = data["timestamp"]
        if not env and "env" in data:
            env = data["env"]

        ops = data.get("result", data)
        for op_name, op_data in ops.items():
            if not isinstance(op_data, dict):
                continue
            if "accuracy" not in op_data:
                continue

            if op_name not in merged:
                merged[op_name] = op_data
            else:
                existing = merged[op_name]
                new_acc = op_data.get("accuracy", {})
                old_acc = existing.get("accuracy", {})

                if (
                    new_acc.get("status") == "Failed"
                    and old_acc.get("status") != "Failed"
                ):
                    merged[op_name] = op_data
                elif old_acc.get("status") == "Failed":
                    pass
                elif new_acc.get("total", 0) > old_acc.get("total", 0):
                    merged[op_name] = op_data

    return merged, env, timestamp


def extract_perf_speedups(op_data):
    """从 performance 数据中提取各 dtype 的 speedup 值"""
    speedups = {}
    perf = op_data.get("performance", {})

    if not isinstance(perf, dict) or not perf:
        return speedups, 0.0

    perf_data = perf.get("data", {})
    if not isinstance(perf_data, dict):
        return speedups, 0.0

    for dtype, d in perf_data.items():
        if isinstance(d, dict):
            sp = d.get("speedup")
            if sp is not None and isinstance(sp, (int, float)):
                speedups[dtype] = float(sp)

    values = [v for v in speedups.values() if v > 0]
    avg = (
        float(np.mean(values))
        if values and np
        else (sum(values) / len(values) if values else 0.0)
    )

    return speedups, avg


def infer_responsibility(reasons):
    """从失败/跳过原因中推断责任人"""
    text = " ".join(reasons).lower()
    tags = []
    if "vllm" in text:
        tags.append("vllm")
    if "cuda" in text:
        tags.append("cuda")
    if "hopper" in text:
        tags.append("hopper")
    if "torch_musa" in text or "torch" in text:
        tags.append("torch")
    if "flaggems" in text:
        tags.append("FlagGems")
    if "transformer engine" in text or "transformerengine" in text:
        tags.append("transformer_engine")
    if "cpu mode" in text or "unsupported in cpu" in text:
        tags.append("cpu")
    if "musa" in text and "torch" not in text:
        tags.append("musa")
    return " & ".join(tags) if tags else ""


def format_problem(op_name, op_data):
    """格式化存在问题描述（简短）"""
    acc = op_data.get("accuracy", {})
    details = acc.get("details", {})

    reasons = []

    if isinstance(details, dict):
        for category in ("failed", "skipped"):
            cat_data = details.get(category, {})
            if isinstance(cat_data, dict):
                for reason in cat_data.keys():
                    # 取第一行并截断到 80 字符
                    first_line = reason.split("\n")[0].strip()
                    if len(first_line) > 80:
                        first_line = first_line[:77] + "..."
                    reasons.append(first_line)
            elif isinstance(cat_data, str):
                first_line = cat_data.split("\n")[0].strip()
                if len(first_line) > 80:
                    first_line = first_line[:77] + "..."
                reasons.append(first_line)

    status = acc.get("status", "")

    if reasons:
        return "; ".join(reasons[:3])
    elif acc.get("total", 0) == 0:
        return "无精度测试用例"
    elif status == "Error":
        return "测试执行错误 (Error)"
    elif acc.get("failed", 0) > 0:
        return "精度测试失败"
    elif acc.get("skipped", 0) == acc.get("total", 0) and acc.get("total", 0) > 0:
        return "全部用例跳过"
    return ""


def analyze_data(results):
    """分析所有算子测试结果"""
    ops_list = []

    for op_name, op_data in results.items():
        acc = op_data.get("accuracy", {})
        total = acc.get("total", 0)
        passed = acc.get("passed", 0)
        failed = acc.get("failed", 0)
        skipped = acc.get("skipped", 0)
        status = acc.get("status", "Unknown")
        details = acc.get("details", {})
        customized = op_data.get("customized", False)

        speedups, avg_speedup = extract_perf_speedups(op_data)
        has_perf = avg_speedup > 0

        if total == 0:
            category = "no_tests"
        elif failed > 0:
            category = "failed"
        elif skipped == total and total > 0:
            category = "skipped"
        elif status == "Error":
            category = "error"
        elif status == "Passed" or failed == 0:
            category = "passed"
        else:
            category = "unknown"

        problem = format_problem(op_name, op_data)

        reasons_for_infer = []
        if isinstance(details, dict):
            for cat in ("failed", "skipped"):
                cat_data = details.get(cat, {})
                if isinstance(cat_data, dict):
                    reasons_for_infer.extend(cat_data.keys())

        responsibility = infer_responsibility(reasons_for_infer)

        op_info = {
            "name": op_name,
            "customized": customized,
            "total": total,
            "passed": passed,
            "failed": failed,
            "skipped": skipped,
            "status": status,
            "category": category,
            "avg_speedup": avg_speedup,
            "has_perf": has_perf,
            "speedups": speedups,
            "details": details,
            "problem": problem,
            "responsibility": responsibility,
        }
        ops_list.append(op_info)

    ops_list.sort(key=lambda x: x["name"])
    return ops_list


def generate_excel(ops_list, output_path):
    """生成 summary.xlsx (OpList + SpeedUp 两个 sheet)"""
    wb = openpyxl.Workbook()
    default_font = Font(name="Microsoft YaHei", size=10)
    header_font = Font(name="Microsoft YaHei", size=10, bold=True)
    blue_fill = PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid")
    header_fill = PatternFill(
        start_color="CCCCCC", end_color="CCCCCC", fill_type="solid"
    )
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    # ---- Sheet 1: OpList ----
    ws1 = wb.active
    ws1.title = "OpList"
    op_headers = ["算子名", "存在问题", "责任人"]
    for col_idx, h in enumerate(op_headers, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_align

    for row_idx, op in enumerate(ops_list, 2):
        ws1.cell(row=row_idx, column=1, value=op["name"]).border = thin_border
        ws1.cell(row=row_idx, column=2, value=op["problem"] or "").border = thin_border
        ws1.cell(
            row=row_idx, column=3, value=op["responsibility"] or ""
        ).border = thin_border

        if op["category"] in ("failed", "skipped", "error", "no_tests"):
            ws1.cell(row=row_idx, column=1).fill = blue_fill

    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 60
    ws1.column_dimensions["C"].width = 20

    # ---- Sheet 2: SpeedUp ----
    ws2 = wb.create_sheet("SpeedUp")
    dtype_cols = ["bool", "int32", "fp32", "fp16", "bf16", "int16", "cf64"]
    speed_headers = [
        "算子名",
        "精度测例总数",
        "精度测例通过数",
        "精度测例失败数",
        "精度测例skip数",
        "性能加速比",
    ] + dtype_cols

    for col_idx, h in enumerate(speed_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = center_align

    for row_idx, op in enumerate(ops_list, 2):
        ws2.cell(row=row_idx, column=1, value=op["name"]).border = thin_border
        ws2.cell(row=row_idx, column=2, value=op["total"]).border = thin_border
        ws2.cell(row=row_idx, column=3, value=op["passed"]).border = thin_border
        ws2.cell(row=row_idx, column=4, value=op["failed"]).border = thin_border
        ws2.cell(row=row_idx, column=5, value=op["skipped"]).border = thin_border

        if op["has_perf"]:
            ws2.cell(
                row=row_idx, column=6, value=round(op["avg_speedup"], 6)
            ).border = thin_border
        else:
            ws2.cell(row=row_idx, column=6).border = thin_border

        for dtype_idx, dtype_key in enumerate(dtype_cols):
            col = 7 + dtype_idx
            val = op["speedups"].get(dtype_key)
            if val is not None and val > 0:
                ws2.cell(
                    row=row_idx, column=col, value=round(val, 6)
                ).border = thin_border
            else:
                ws2.cell(row=row_idx, column=col).border = thin_border

    ws2.column_dimensions["A"].width = 30
    for col_idx in range(2, len(speed_headers) + 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 16

    wb.save(output_path)
    print(f"  Excel 已生成: {output_path}")


def generate_html(ops_list, env, timestamp, folder_name):
    """生成 HTML 报告"""
    import html as html_module

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total = len(ops_list)

    passed_ops = [op for op in ops_list if op["category"] == "passed"]
    failed_ops = [op for op in ops_list if op["category"] == "failed"]
    skipped_ops = [op for op in ops_list if op["category"] == "skipped"]
    error_ops = [op for op in ops_list if op["category"] == "error"]
    no_test_ops = [op for op in ops_list if op["category"] == "no_tests"]
    accuracy_passed = len(passed_ops)
    accuracy_failed = len(failed_ops)
    accuracy_skipped = len(skipped_ops)
    accuracy_error = len(error_ops)
    accuracy_no_test = len(no_test_ops)

    perf_passed_ops = [
        op for op in ops_list if op["has_perf"] and op["category"] == "passed"
    ]
    perf_and_acc_pass = len(perf_passed_ops)

    speedups_list = [op["avg_speedup"] for op in ops_list if op["has_perf"]]

    if speedups_list:
        arr = np.array(speedups_list) if np else speedups_list
        median = (
            float(np.median(arr))
            if np
            else sorted(speedups_list)[len(speedups_list) // 2]
        )
        mean = float(np.mean(arr)) if np else sum(speedups_list) / len(speedups_list)
        min_s = float(np.min(arr)) if np else min(speedups_list)
        max_s = float(np.max(arr)) if np else max(speedups_list)

        below_08 = sum(1 for s in speedups_list if s < 0.8)
        between_08_1 = sum(1 for s in speedups_list if 0.8 <= s <= 1.0)
        above_1 = sum(1 for s in speedups_list if s > 1.0)
        total_speedups = len(speedups_list)
    else:
        median = mean = min_s = max_s = 0.0
        below_08 = between_08_1 = above_1 = total_speedups = 0

    # 加速比分布百分比
    if total_speedups > 0:
        pct_below = below_08 / total_speedups * 100
        pct_between = between_08_1 / total_speedups * 100
        pct_above = above_1 / total_speedups * 100
    else:
        pct_below = pct_between = pct_above = 0

    # 图表数据: (func_name, speedup)
    chart_data = json.dumps(
        [[op["name"], round(op["avg_speedup"], 6)] for op in ops_list if op["has_perf"]]
    )

    # 慢算子 (speedup < 0.8)
    slow_ops = sorted(
        [op for op in ops_list if op["has_perf"] and op["avg_speedup"] < 0.8],
        key=lambda x: x["avg_speedup"],
    )
    fast_ops = sorted(
        [op for op in ops_list if op["has_perf"] and op["avg_speedup"] > 2.0],
        key=lambda x: x["avg_speedup"],
        reverse=True,
    )

    slow_rows = ""
    for op in slow_ops:
        slow_rows += f'<tr><td>{op["name"]}</td><td><span class="badge badge-danger">{op["avg_speedup"]:.4f}</span></td></tr>\n'

    fast_rows = ""
    for op in fast_ops:
        fast_rows += f'<tr><td>{op["name"]}</td><td><span class="badge badge-success">{op["avg_speedup"]:.4f}</span></td></tr>\n'

    # 失败/跳过算子详情
    problem_ops = failed_ops + skipped_ops + error_ops + no_test_ops
    problem_ops.sort(key=lambda x: x["name"])

    problem_cards = ""
    for op in problem_ops:
        name = op["name"]
        details = op["details"]

        detail_html = ""
        if isinstance(details, dict):
            for category, items in details.items():
                if isinstance(items, dict):
                    for reason, cases in items.items():
                        safe_reason = html_module.escape(reason[:200])
                        case_count = len(cases) if isinstance(cases, list) else 1
                        detail_html += (
                            f'<div class="detail-reason"><strong>{category}:</strong> '
                            f'{safe_reason}<br><span class="case-count">({case_count} 个用例)</span></div>'
                        )
                elif isinstance(items, str):
                    detail_html += (
                        f'<div class="detail-reason"><strong>{category}:</strong> '
                        f"{html_module.escape(items[:200])}</div>"
                    )

        if op["category"] == "failed":
            badge_class = "badge-danger"
            status_label = f'失败({op["failed"]})'
            card_class = ""
        elif op["category"] == "skipped":
            badge_class = "badge-warning"
            status_label = f'全部跳过({op["skipped"]})'
            card_class = " skipped"
        elif op["category"] == "no_tests":
            badge_class = "badge-warning"
            status_label = "无精度用例"
            card_class = " skipped"
        else:
            badge_class = "badge-danger"
            status_label = "Error"
            card_class = ""

        problem_cards += f"""
        <div class="fail-card{card_class}">
            <div class="fail-header">
                <span class="fail-name">{name}</span>
                <span class="badge {badge_class}">{status_label}</span>
                <span class="custom-badge">{'自定义' if op.get("customized") else '通用'}</span>
            </div>
            <div class="fail-stats">
                总计 {op["total"]} | 通过 {op["passed"]} | 失败 {op["failed"]} | 跳过 {op["skipped"]}
            </div>
            <div class="fail-details">{detail_html if detail_html else '<span class="no-detail">无详细信息</span>'}</div>
        </div>"""

    # 通过算子摘要
    passed_summary = ""
    if len(passed_ops) <= 50:
        passed_summary = ", ".join(op["name"] for op in passed_ops)
    else:
        passed_summary = f'{", ".join(op["name"] for op in passed_ops[:50])} ... 等共 {len(passed_ops)} 个算子'

    # 环境信息
    env_html = ""
    if env:
        device = env.get("flag_gems", {}).get("device", "N/A")
        vendor = env.get("flag_gems", {}).get("vendor", "N/A")
        fg_version = env.get("flag_gems", {}).get("version", "N/A")
        triton_version = env.get("triton", {}).get("version", "N/A")
        torch_version = env.get("torch", {}).get("version", "N/A")
        python_version = env.get("python", "N/A")
        os_name = f'{env.get("os_name", "")} {env.get("os_release", "")}'.strip()

        env_html = f"""
        <div class="env-grid">
            <div class="env-item-card"><span class="env-key">FlagGems</span><span class="env-val">{fg_version}</span></div>
            <div class="env-item-card"><span class="env-key">Triton</span><span class="env-val">{triton_version}</span></div>
            <div class="env-item-card"><span class="env-key">PyTorch</span><span class="env-val">{torch_version}</span></div>
            <div class="env-item-card"><span class="env-key">Python</span><span class="env-val">{python_version}</span></div>
            <div class="env-item-card"><span class="env-key">Device</span><span class="env-val">{device} ({vendor})</span></div>
            <div class="env-item-card"><span class="env-key">OS</span><span class="env-val">{os_name}</span></div>
        </div>"""

    # 优先/次优先优化
    priority_ops = [
        f"<strong>{op['name']}</strong> ({op['avg_speedup']:.2f})"
        for op in slow_ops[:4]
    ]
    priority_html = "<br>".join(priority_ops) if priority_ops else "无"

    secondary_ops = [
        f"<strong>{op['name']}</strong> ({op['avg_speedup']:.2f})"
        for op in slow_ops[4:8]
    ]
    secondary_html = "<br>".join(secondary_ops) if secondary_ops else "无"

    pass_rate = (accuracy_passed / total * 100) if total > 0 else 0

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>算子测试结果分析报告</title>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, 'Helvetica Neue', Arial, sans-serif;
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            min-height: 100vh;
            padding: 40px 20px;
        }}
        .container {{ max-width: 1200px; margin: 0 auto; }}
        .header {{ text-align: center; color: white; margin-bottom: 40px; }}
        .header h1 {{ font-size: 2.5rem; margin-bottom: 10px; text-shadow: 2px 2px 4px rgba(0,0,0,0.2); }}
        .card {{
            background: white;
            border-radius: 16px;
            box-shadow: 0 10px 40px rgba(0,0,0,0.15);
            margin-bottom: 30px;
            overflow: hidden;
        }}
        .card-header {{
            background: linear-gradient(135deg, #5a67d8 0%, #6b46c1 100%);
            color: white;
            padding: 20px 30px;
            font-size: 1.3rem;
            font-weight: 600;
        }}
        .card-body {{ padding: 30px; }}
        .stats-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
        }}
        .stat-item {{
            background: linear-gradient(135deg, #f6f8fc 0%, #eef2f7 100%);
            border-radius: 12px;
            padding: 25px;
            text-align: center;
            transition: transform 0.3s ease;
        }}
        .stat-item:hover {{ transform: translateY(-5px); }}
        .stat-value {{ font-size: 2.5rem; font-weight: 700; color: #5a67d8; margin-bottom: 8px; }}
        .stat-value.success {{ color: #38a169; }}
        .stat-value.danger {{ color: #e53e3e; }}
        .stat-value.warning {{ color: #d69e2e; }}
        .stat-label {{ color: #718096; font-size: 0.95rem; }}
        table {{ width: 100%; border-collapse: collapse; }}
        th, td {{ padding: 12px 16px; text-align: left; border-bottom: 1px solid #e2e8f0; }}
        th {{ background: #f7fafc; font-weight: 600; color: #4a5568; text-transform: uppercase; font-size: 0.85rem; }}
        tr:hover {{ background: #f7fafc; }}
        .badge {{ display: inline-block; padding: 4px 12px; border-radius: 20px; font-size: 0.85rem; font-weight: 500; }}
        .badge-success {{ background: #c6f6d5; color: #22543d; }}
        .badge-warning {{ background: #feebc8; color: #744210; }}
        .badge-danger {{ background: #fed7d7; color: #742a2a; }}
        .summary-box {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 15px; margin-bottom: 25px; }}
        .summary-item {{ text-align: center; padding: 20px; background: #f7fafc; border-radius: 10px; }}
        .summary-item .value {{ font-size: 1.8rem; font-weight: 700; color: #2d3748; }}
        .summary-item .label {{ font-size: 0.9rem; color: #718096; margin-top: 5px; }}
        .distribution-chart {{ display: flex; height: 40px; border-radius: 8px; overflow: hidden; margin: 20px 0; }}
        .dist-segment {{ display: flex; align-items: center; justify-content: center; color: white; font-weight: 600; font-size: 0.9rem; }}
        .dist-low {{ background: linear-gradient(90deg, #fc8181, #f56565); }}
        .dist-medium {{ background: linear-gradient(90deg, #f6ad55, #ed8936); }}
        .dist-high {{ background: linear-gradient(90deg, #68d391, #48bb78); }}
        .legend {{ display: flex; justify-content: center; gap: 30px; margin-top: 15px; }}
        .legend-item {{ display: flex; align-items: center; gap: 8px; font-size: 0.9rem; color: #4a5568; }}
        .legend-dot {{ width: 12px; height: 12px; border-radius: 50%; }}
        .two-col {{ display: grid; grid-template-columns: 1fr 1fr; gap: 30px; }}
        .section-title {{ font-size: 1.1rem; color: #4a5568; margin-bottom: 20px; padding-bottom: 10px; border-bottom: 2px solid #e2e8f0; }}
        .op-list {{ max-height: 400px; overflow-y: auto; }}
        .fail-card {{
            background: #fff5f5;
            border-left: 4px solid #fc8181;
            padding: 15px 20px;
            border-radius: 8px;
            margin-bottom: 12px;
        }}
        .fail-card.skipped {{ background: #fffff0; border-left-color: #ecc94b; }}
        .fail-header {{ display: flex; align-items: center; gap: 12px; margin-bottom: 8px; }}
        .fail-name {{ font-weight: 600; color: #2d3748; font-size: 1.05rem; }}
        .fail-stats {{ color: #718096; font-size: 0.85rem; margin-bottom: 8px; }}
        .fail-details {{ font-size: 0.82rem; max-height: 200px; overflow-y: auto; }}
        .detail-reason {{ margin-bottom: 4px; color: #742a2a; font-family: monospace; font-size: 0.78rem; }}
        .case-count {{ color: #718096; font-family: sans-serif; }}
        .custom-badge {{
            display: inline-block;
            padding: 2px 8px;
            border-radius: 4px;
            font-size: 0.75rem;
            background: #e2e8f0;
            color: #4a5568;
        }}
        .no-detail {{ color: #a0aec0; font-style: italic; }}
        .env-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
            gap: 12px;
        }}
        .env-item-card {{
            background: #f7fafc;
            border-radius: 8px;
            padding: 12px;
            text-align: center;
        }}
        .env-item-card .env-key {{ display: block; font-size: 0.8rem; color: #718096; margin-bottom: 4px; }}
        .env-item-card .env-val {{ display: block; font-weight: 600; color: #2d3748; font-size: 0.9rem; }}
        .passed-ops {{ max-height: 300px; overflow-y: auto; line-height: 1.8; }}
        .search-box {{
            width: 100%;
            padding: 10px 16px;
            border: 1px solid #e2e8f0;
            border-radius: 8px;
            font-size: 0.95rem;
            margin-bottom: 15px;
        }}
        .search-box:focus {{ outline: none; border-color: #5a67d8; box-shadow: 0 0 0 3px rgba(90,103,216,0.2); }}
        .footer {{ text-align: center; color: rgba(255,255,255,0.7); margin-top: 30px; font-size: 0.9rem; }}
        .chart-container {{ position: relative; }}
        @media (max-width: 768px) {{
            .two-col {{ grid-template-columns: 1fr; }}
            .summary-box {{ grid-template-columns: repeat(2, 1fr); }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>算子测试结果分析报告</h1>
            <div style="margin-top: 20px; font-size: 1.1rem;">
                <span>{folder_name}</span> | <span>{total} 个算子</span>
            </div>
        </div>

        <div class="card">
            <div class="card-header">1. 概览</div>
            <div class="card-body">
                <div class="stats-grid">
                    <div class="stat-item">
                        <div class="stat-value">{total}</div>
                        <div class="stat-label">总算子数量</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value success">{accuracy_passed}</div>
                        <div class="stat-label">精度通过数</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value success">{perf_and_acc_pass}</div>
                        <div class="stat-label">精度与性能均通过</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value danger">{accuracy_failed}</div>
                        <div class="stat-label">精度失败数</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value warning">{accuracy_no_test}</div>
                        <div class="stat-label">无精度用例</div>
                    </div>
                    <div class="stat-item">
                        <div class="stat-value warning">{accuracy_skipped}</div>
                        <div class="stat-label">跳过数</div>
                    </div>
                </div>
                {f'<div style="margin-top: 15px; color: #718096; font-size: 0.9rem;">Error: {accuracy_error} 个</div>' if accuracy_error > 0 else ''}
            </div>
        </div>

        <div class="card">
            <div class="card-header">2. 加速比统计</div>
            <div class="card-body">
                <div class="summary-box">
                    <div class="summary-item">
                        <div class="value">{median:.4f}</div>
                        <div class="label">中位数</div>
                    </div>
                    <div class="summary-item">
                        <div class="value">{mean:.4f}</div>
                        <div class="label">平均值</div>
                    </div>
                    <div class="summary-item">
                        <div class="value">{min_s:.4f}</div>
                        <div class="label">最小值</div>
                    </div>
                    <div class="summary-item">
                        <div class="value">{max_s:.4f}</div>
                        <div class="label">最大值</div>
                    </div>
                </div>

                <h3 class="section-title">加速比分布</h3>
                <div class="distribution-chart">
                    <div class="dist-segment dist-low" style="flex: {pct_below};">{pct_below:.1f}%</div>
                    <div class="dist-segment dist-medium" style="flex: {pct_between};">{pct_between:.1f}%</div>
                    <div class="dist-segment dist-high" style="flex: {pct_above};">{pct_above:.1f}%</div>
                </div>
                <div class="legend">
                    <div class="legend-item"><div class="legend-dot" style="background: #f56565;"></div><span>&lt; 0.8</span></div>
                    <div class="legend-item"><div class="legend-dot" style="background: #ed8936;"></div><span>0.8 ~ 1.0</span></div>
                    <div class="legend-item"><div class="legend-dot" style="background: #48bb78;"></div><span>&gt; 1.0</span></div>
                </div>

                <table style="margin-top: 30px;">
                    <thead>
                        <tr><th>区间</th><th>数量</th><th>占比</th></tr>
                    </thead>
                    <tbody>
                        <tr><td><span class="badge badge-danger">&lt; 0.8</span></td><td>{below_08}</td><td>{pct_below:.2f}%</td></tr>
                        <tr><td><span class="badge badge-warning">0.8 ~ 1.0</span></td><td>{between_08_1}</td><td>{pct_between:.2f}%</td></tr>
                        <tr><td><span class="badge badge-success">&gt; 1.0</span></td><td>{above_1}</td><td>{pct_above:.2f}%</td></tr>
                    </tbody>
                </table>
            </div>
        </div>

        <div class="card">
            <div class="card-header">3. 加速比柱状图</div>
            <div class="card-body">
                <div style="display: flex; flex-wrap: wrap; gap: 15px; margin-bottom: 20px; align-items: center;">
                    <div>
                        <label style="font-size: 0.9rem; color: #4a5568; margin-right: 8px;">筛选:</label>
                        <select id="filterRange" onchange="updateChart()" style="padding: 8px 12px; border: 1px solid #e2e8f0; border-radius: 6px;">
                            <option value="all">全部</option>
                            <option value="below08">&lt; 0.8</option>
                            <option value="between">0.8 ~ 1.0</option>
                            <option value="above1">&gt; 1.0</option>
                        </select>
                    </div>
                    <div>
                        <label style="font-size: 0.9rem; color: #4a5568; margin-right: 8px;">排序:</label>
                        <select id="sortOrder" onchange="updateChart()" style="padding: 8px 12px; border: 1px solid #e2e8f0; border-radius: 6px;">
                            <option value="name">按名称</option>
                            <option value="asc">加速比升序</option>
                            <option value="desc">加速比降序</option>
                        </select>
                    </div>
                    <div>
                        <label style="font-size: 0.9rem; color: #4a5568; margin-right: 8px;">Y轴上限:</label>
                        <input type="number" id="yAxisMax" value="3" min="1" step="0.5" onchange="updateChart()" style="padding: 8px 12px; border: 1px solid #e2e8f0; border-radius: 6px; width: 80px;">
                    </div>
                    <div style="flex: 1; min-width: 200px;">
                        <input type="text" id="searchBox" placeholder="搜索算子名..." oninput="updateChart()" style="padding: 8px 12px; border: 1px solid #e2e8f0; border-radius: 6px; width: 100%;">
                    </div>
                </div>
                <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px;">
                    <span id="chartInfo" style="font-size: 0.9rem; color: #718096;"></span>
                    <div>
                        <button onclick="prevPage()" style="padding: 6px 12px; border: 1px solid #e2e8f0; border-radius: 6px; cursor: pointer; margin-right: 5px;">&lt; 上一页</button>
                        <span id="pageInfo" style="font-size: 0.9rem; color: #4a5568; margin: 0 10px;"></span>
                        <button onclick="nextPage()" style="padding: 6px 12px; border: 1px solid #e2e8f0; border-radius: 6px; cursor: pointer;">下一页 &gt;</button>
                    </div>
                </div>
                <div class="chart-container" style="height: 400px;">
                    <canvas id="speedupChart"></canvas>
                </div>
            </div>
        </div>

        <div class="two-col">
            <div class="card">
                <div class="card-header">4. 需关注算子（加速比 &lt; 0.8）</div>
                <div class="card-body">
                    <div class="op-list">
                        <table>
                            <thead><tr><th>算子名</th><th>加速比</th></tr></thead>
                            <tbody>{slow_rows if slow_rows else '<tr><td colspan="2" style="text-align:center;color:#718096;">无</td></tr>'}</tbody>
                        </table>
                    </div>
                </div>
            </div>

            <div class="card">
                <div class="card-header">5. 高性能算子（加速比 &gt; 2.0）</div>
                <div class="card-body">
                    <div class="op-list">
                        <table>
                            <thead><tr><th>算子名</th><th>加速比</th></tr></thead>
                            <tbody>{fast_rows if fast_rows else '<tr><td colspan="2" style="text-align:center;color:#718096;">无</td></tr>'}</tbody>
                        </table>
                    </div>
                </div>
            </div>
        </div>

        <div class="card">
            <div class="card-header">6. 精度测试失败 / 跳过 / 无精度用例 / Error 算子</div>
            <div class="card-body">
                <input type="text" class="search-box" placeholder="搜索算子名..." oninput="filterFails(this.value)">
                <div id="failContainer">
                    {problem_cards if problem_cards else '<p style="text-align: center; color: #38a169; font-size: 1.1rem;">无失败或跳过算子，全部通过！</p>'}
                </div>
            </div>
        </div>

        <div class="card">
            <div class="card-header">7. 精度测试通过算子</div>
            <div class="card-body">
                <div class="passed-ops">{passed_summary}</div>
            </div>
        </div>

        <div class="card">
            <div class="card-header">8. 测试环境</div>
            <div class="card-body">
                {env_html if env_html else '<p style="color: #a0aec0;">未找到环境信息</p>'}
                {'<p style="color: #718096; font-size: 0.9rem; margin-top: 10px;">测试时间: ' + timestamp + '</p>' if timestamp else ''}
            </div>
        </div>

        <div class="card">
            <div class="card-header">9. 优化建议</div>
            <div class="card-body">
                <div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(300px, 1fr)); gap: 20px;">
                    <div style="background: #fff5f5; border-left: 4px solid #fc8181; padding: 20px; border-radius: 8px;">
                        <h4 style="color: #c53030; margin-bottom: 10px;">优先优化</h4>
                        <p style="color: #742a2a; font-size: 0.95rem;">{priority_html}</p>
                    </div>
                    <div style="background: #fffff0; border-left: 4px solid #ecc94b; padding: 20px; border-radius: 8px;">
                        <h4 style="color: #975a16; margin-bottom: 10px;">次优先优化</h4>
                        <p style="color: #744210; font-size: 0.95rem;">{secondary_html}</p>
                    </div>
                    <div style="background: #f0fff4; border-left: 4px solid #68d391; padding: 20px; border-radius: 8px;">
                        <h4 style="color: #276749; margin-bottom: 10px;">优势保持</h4>
                        <p style="color: #22543d; font-size: 0.95rem;">
                            通过率 {pass_rate:.1f}%，其中 {len(fast_ops)} 个算子加速比超过 2 倍。
                        </p>
                    </div>
                </div>
            </div>
        </div>

        <div class="footer">
            报告生成时间: {now}
        </div>
    </div>

    <script>
        const allData = {chart_data};
        let filteredData = [];
        let currentPage = 0;
        const pageSize = 50;
        let chart = null;

        function filterAndSortData() {{
            const filterRange = document.getElementById('filterRange').value;
            const sortOrder = document.getElementById('sortOrder').value;
            const searchText = document.getElementById('searchBox').value.toLowerCase();

            filteredData = allData.filter(item => {{
                const [name, speedup] = item;
                if (searchText && !name.toLowerCase().includes(searchText)) return false;
                if (filterRange === 'below08' && speedup >= 0.8) return false;
                if (filterRange === 'between' && (speedup < 0.8 || speedup > 1.0)) return false;
                if (filterRange === 'above1' && speedup <= 1.0) return false;
                return true;
            }});

            if (sortOrder === 'asc') {{
                filteredData.sort((a, b) => a[1] - b[1]);
            }} else if (sortOrder === 'desc') {{
                filteredData.sort((a, b) => b[1] - a[1]);
            }} else {{
                filteredData.sort((a, b) => a[0].localeCompare(b[0]));
            }}
        }}

        function updateChart() {{
            filterAndSortData();
            currentPage = 0;
            renderChart();
        }}

        function renderChart() {{
            const yAxisMax = parseFloat(document.getElementById('yAxisMax').value) || 3;
            const start = currentPage * pageSize;
            const end = Math.min(start + pageSize, filteredData.length);
            const pageData = filteredData.slice(start, end);

            const labels = pageData.map(item => item[0]);
            const values = pageData.map(item => item[1]);
            const displayValues = values.map(v => Math.min(v, yAxisMax));
            const colors = values.map(v => {{
                if (v < 0.8) return 'rgba(245, 101, 101, 0.8)';
                if (v <= 1.0) return 'rgba(237, 137, 54, 0.8)';
                return 'rgba(72, 187, 120, 0.8)';
            }});

            document.getElementById('chartInfo').innerText = `共 ${{filteredData.length}} 个算子，当前显示 ${{start + 1}}-${{end}}`;
            document.getElementById('pageInfo').innerText = `${{currentPage + 1}} / ${{Math.ceil(filteredData.length / pageSize) || 1}}`;

            if (chart) chart.destroy();

            const ctx = document.getElementById('speedupChart').getContext('2d');
            chart = new Chart(ctx, {{
                type: 'bar',
                data: {{
                    labels: labels,
                    datasets: [{{
                        label: '加速比',
                        data: displayValues,
                        backgroundColor: colors,
                        borderColor: colors.map(c => c.replace('0.8', '1')),
                        borderWidth: 1
                    }}]
                }},
                options: {{
                    responsive: true,
                    maintainAspectRatio: false,
                    plugins: {{
                        tooltip: {{
                            callbacks: {{
                                label: function(context) {{
                                    const realValue = values[context.dataIndex];
                                    if (realValue > yAxisMax) {{
                                        return `加速比: ${{realValue.toFixed(4)}} (截断显示)`;
                                    }}
                                    return `加速比: ${{realValue.toFixed(4)}}`;
                                }}
                            }}
                        }},
                        legend: {{ display: false }}
                    }},
                    scales: {{
                        y: {{
                            beginAtZero: true,
                            max: yAxisMax,
                            title: {{ display: true, text: '加速比' }}
                        }},
                        x: {{
                            ticks: {{
                                maxRotation: 45,
                                minRotation: 45,
                                font: {{ size: 10 }}
                            }}
                        }}
                    }}
                }}
            }});
        }}

        function prevPage() {{
            if (currentPage > 0) {{
                currentPage--;
                renderChart();
            }}
        }}

        function nextPage() {{
            if ((currentPage + 1) * pageSize < filteredData.length) {{
                currentPage++;
                renderChart();
            }}
        }}

        function filterFails(query) {{
            const cards = document.querySelectorAll('.fail-card');
            const q = query.toLowerCase();
            cards.forEach(card => {{
                const name = card.querySelector('.fail-name').innerText.toLowerCase();
                card.style.display = name.includes(q) ? '' : 'none';
            }});
        }}

        document.addEventListener('DOMContentLoaded', updateChart);
    </script>
</body>
</html>"""

    return html


def main():
    if len(sys.argv) < 2:
        print("用法: python analyze_ops_gen_html_summary_from_json.py <结果文件夹路径>")
        print("示例: python analyze_ops_gen_html_summary_from_json.py result_20260512")
        sys.exit(1)

    folder = Path(sys.argv[1])
    if not folder.exists():
        print(f"错误: 文件夹不存在: {folder}")
        sys.exit(1)

    print(f"分析目录: {folder}")

    # 加载数据
    results, env, timestamp = load_results(folder)
    print(f"共加载 {len(results)} 个算子")

    # 分析
    ops_list = analyze_data(results)
    passed = sum(1 for op in ops_list if op["category"] == "passed")
    failed = sum(1 for op in ops_list if op["category"] == "failed")
    skipped = sum(1 for op in ops_list if op["category"] == "skipped")
    error = sum(1 for op in ops_list if op["category"] == "error")
    no_test = sum(1 for op in ops_list if op["category"] == "no_tests")
    print(
        f"通过: {passed}, 失败: {failed}, 跳过: {skipped}, 无精度用例: {no_test}, Error: {error}"
    )

    # 生成 Excel
    xlsx_path = folder / "summary.xlsx"
    generate_excel(ops_list, xlsx_path)

    # 生成 HTML
    html = generate_html(ops_list, env, timestamp, folder.name)
    html_path = folder / "report.html"
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"  HTML 已生成: {html_path}")

    # 摘要
    perf_ops = [op for op in ops_list if op["has_perf"]]
    if perf_ops:
        sps = [op["avg_speedup"] for op in perf_ops]
        arr = np.array(sps) if np else sps
        median = float(np.median(arr)) if np else sorted(sps)[len(sps) // 2]
        mean = float(np.mean(arr)) if np else sum(sps) / len(sps)
        print(f"\n===== 分析摘要 =====")
        print(f"总算子数: {len(ops_list)}")
        print(f"精度通过: {passed}")
        print(f"加速比中位数: {median:.4f}")
        print(f"加速比平均值: {mean:.4f}")


if __name__ == "__main__":
    main()
